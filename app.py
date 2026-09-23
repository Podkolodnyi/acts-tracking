
from flask_cors import CORS
from functools import partial
import csv
import io
import os
import sqlite3
import urllib.request
from contextlib import contextmanager
from datetime import datetime
from io import BytesIO

import requests
from dotenv import load_dotenv
from flask import (
    Flask,
    flash,
    jsonify,
    redirect,
    render_template,
    request,
    send_file,
    session,
    url_for,
)
from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter
from werkzeug.security import check_password_hash, generate_password_hash

BASE_DIR = os.path.dirname(os.path.abspath(__file__))
load_dotenv(os.path.join(BASE_DIR, ".env"))

INTRASERVICE_BASE_URL = os.environ["INTRASERVICE_BASE_URL"].rstrip("/")
INTRASERVICE_AUTH = (
    os.environ["INTRASERVICE_API_LOGIN"],
    os.environ["INTRASERVICE_API_PASSWORD"],
)

DATABASE = os.path.join(BASE_DIR, "repair_acts.db")
SHEET_URL = (
    "https://docs.google.com/spreadsheets/d/"
    "1wq-TfC-wGvyZTFTQLSbsxles7KJ3Owj3KYXeiAcoeP8/"
    "export?format=csv&gid=467885286"
)

# One-time seed data for the `engineers` table (see init_db). Editing this
# after the first run has no effect — it only backfills an empty table.
SEED_ENGINEERS = {
    "podkolodny": {"name": "Подколодный", "code": "1", "is_admin": True},
    "dzyuba": {"name": "Дзюба", "code": "2", "is_admin": False},
    "izyurov": {"name": "Изъюров", "code": "5", "is_admin": False},
    "ozerov": {"name": "Озеров", "code": "9", "is_admin": False},
}

NUMBER_TYPES = {
    "standard": {"name": "Обычный акт"},
    "thermo": {"name": "Узел терморегистрации"},
}

app = Flask(__name__)
app.config["SECRET_KEY"] = os.environ.get(
    "FLASK_SECRET_KEY",
    "change-this-before-network-use",
)
# Enable CORS for React frontend (with credentials, so the session cookie works)
FRONTEND_ORIGIN = os.environ.get("FRONTEND_ORIGIN", "http://localhost:5173")
CORS(app, supports_credentials=True, origins=[FRONTEND_ORIGIN])


@contextmanager
def db():
    connection = sqlite3.connect(DATABASE)
    connection.row_factory = sqlite3.Row
    try:
        yield connection
        connection.commit()
    finally:
        connection.close()


def clean(value):
    return " ".join((value or "").strip().split())


def clean_multiline(value):
    lines = []
    for line in (value or "").splitlines():
        line = " ".join(line.strip().split())
        if line:
            lines.append(line)
    return "\n".join(lines)


def normalize_serial(value):
    return clean(value).upper().replace(" ", "")


def field(row, *names):
    lower = {clean(key).lower(): value for key, value in row.items()}
    for name in names:
        value = lower.get(name.lower())
        if value is not None:
            return clean(value)
    return ""


def search_intraservice(query):
    query = clean(query)
    if not query:
        return []

    response = requests.get(
        f"{INTRASERVICE_BASE_URL}/api/task",
        params={"pagesize": 20, "archive": "true", "search": query},
        auth=INTRASERVICE_AUTH,
        headers={"Accept": "application/json"},
        timeout=20,
    )
    response.raise_for_status()
    payload = response.json()

    return [
        {
            "id": task.get("Id"),
            "name": task.get("Name") or "",
            "description": task.get("Description") or "",
            "created": task.get("Created") or "",
            "deadline": task.get("Deadline") or "",
        }
        for task in payload.get("Tasks", [])
    ]


def init_db():
    with db() as con:
        con.executescript(
            """
            CREATE TABLE IF NOT EXISTS devices (
                id INTEGER PRIMARY KEY AUTOINCREMENT,
                source TEXT NOT NULL,
                source_key TEXT NOT NULL,
                sap_id TEXT,
                mvz TEXT,
                customer_name TEXT NOT NULL,
                object_format TEXT,
                address TEXT,
                model TEXT,
                serial_number TEXT NOT NULL,
                serial_normalized TEXT NOT NULL,
                imported_at TEXT NOT NULL,
                UNIQUE(source, source_key)
            );

            CREATE INDEX IF NOT EXISTS idx_devices_serial
                ON devices(serial_normalized);
            CREATE INDEX IF NOT EXISTS idx_devices_customer
                ON devices(customer_name);

            CREATE TABLE IF NOT EXISTS acts (
                id INTEGER PRIMARY KEY AUTOINCREMENT,
                act_number TEXT UNIQUE,
                number_type TEXT NOT NULL DEFAULT 'standard',
                engineer_key TEXT NOT NULL,
                engineer_name TEXT NOT NULL,
                engineer_code TEXT NOT NULL,
                created_at TEXT NOT NULL,
                updated_at TEXT NOT NULL,
                status TEXT NOT NULL DEFAULT 'draft',
                customer_name TEXT,
                customer_representative TEXT,
                device_model TEXT,
                serial_number TEXT,
                printeco_label TEXT,
                device_type TEXT,
                comment_label TEXT,
                address TEXT,
                phone TEXT,
                fault TEXT,
                counter_bw TEXT,
                counter_color TEXT,
                service_kind TEXT,
                diagnostics_result TEXT,
                works_text TEXT,
                materials_text TEXT,
                work_date TEXT,
                start_time TEXT,
                end_time TEXT,
                device_condition TEXT,
                customer_signatory TEXT,
                source_device_id INTEGER,
                intraservice_task_id TEXT,
                FOREIGN KEY(source_device_id) REFERENCES devices(id)
            );

            CREATE INDEX IF NOT EXISTS idx_acts_serial
                ON acts(serial_number);
            CREATE INDEX IF NOT EXISTS idx_acts_number
                ON acts(act_number);
            CREATE INDEX IF NOT EXISTS idx_acts_engineer
                ON acts(engineer_key);

            CREATE TABLE IF NOT EXISTS engineers (
                key TEXT PRIMARY KEY,
                name TEXT NOT NULL,
                code TEXT NOT NULL UNIQUE,
                password_hash TEXT,
                is_admin INTEGER NOT NULL DEFAULT 0,
                status TEXT NOT NULL DEFAULT 'active',
                created_at TEXT NOT NULL
            );
            """
        )

        columns = {
            row["name"]
            for row in con.execute("PRAGMA table_info(acts)").fetchall()
        }
        if "number_type" not in columns:
            con.execute(
                "ALTER TABLE acts "
                "ADD COLUMN number_type TEXT NOT NULL DEFAULT 'standard'"
            )

        engineer_columns = {
            row["name"]
            for row in con.execute("PRAGMA table_info(engineers)").fetchall()
        }
        if "first_name" not in engineer_columns:
            con.execute("ALTER TABLE engineers ADD COLUMN first_name TEXT")

        has_engineers = con.execute(
            "SELECT 1 FROM engineers LIMIT 1"
        ).fetchone()
        if not has_engineers:
            now = datetime.now().isoformat(timespec="seconds")
            con.executemany(
                "INSERT INTO engineers "
                "(key, name, code, password_hash, is_admin, status, created_at) "
                "VALUES (?, ?, ?, NULL, ?, 'active', ?)",
                [
                    (key, value["name"], value["code"], int(value["is_admin"]), now)
                    for key, value in SEED_ENGINEERS.items()
                ],
            )


def get_engineers_dict(con=None):
    if con is None:
        with db() as fresh_con:
            return get_engineers_dict(fresh_con)

    rows = con.execute(
        "SELECT key, name, code, is_admin, status FROM engineers"
    ).fetchall()
    return {
        row["key"]: {
            "name": row["name"],
            "code": row["code"],
            "is_admin": bool(row["is_admin"]),
            "status": row["status"],
        }
        for row in rows
    }


def get_engineer():
    key = session.get("engineer_key")
    if not key:
        return None

    with db() as con:
        row = con.execute(
            "SELECT name, first_name, code, is_admin FROM engineers "
            "WHERE key = ? AND status = 'active'",
            (key,),
        ).fetchone()

    if not row:
        return None

    return {
        "name": row["name"],
        "first_name": row["first_name"] or "",
        "code": row["code"],
        "is_admin": bool(row["is_admin"]),
    }


def require_engineer():
    if not get_engineer():
        flash("Сначала выберите инженера.", "warning")
        return False
    return True


def require_admin():
    engineer = get_engineer()
    return bool(engineer and engineer["is_admin"])


CYRILLIC_TO_LATIN = {
    "а": "a", "б": "b", "в": "v", "г": "g", "д": "d", "е": "e", "ё": "e",
    "ж": "zh", "з": "z", "и": "i", "й": "y", "к": "k", "л": "l", "м": "m",
    "н": "n", "о": "o", "п": "p", "р": "r", "с": "s", "т": "t", "у": "u",
    "ф": "f", "х": "kh", "ц": "ts", "ч": "ch", "ш": "sh", "щ": "shch",
    "ъ": "", "ы": "y", "ь": "", "э": "e", "ю": "yu", "я": "ya",
}


def transliterate(text):
    return "".join(
        CYRILLIC_TO_LATIN.get(ch, ch if ch.isalnum() else "")
        for ch in text.lower()
    )


def generate_engineer_key(con, first_name, last_name):
    def is_free(candidate):
        return not con.execute(
            "SELECT 1 FROM engineers WHERE key = ?", (candidate,)
        ).fetchone()

    base = transliterate(last_name)
    if base and is_free(base):
        return base

    combined = f"{base}_{transliterate(first_name)}".strip("_")
    if combined and is_free(combined):
        return combined

    fallback = combined or base or "engineer"
    suffix = 2
    while not is_free(f"{fallback}{suffix}"):
        suffix += 1
    return f"{fallback}{suffix}"


def generate_engineer_code(con):
    used_codes = {
        int(row["code"])
        for row in con.execute("SELECT code FROM engineers").fetchall()
        if row["code"].isdigit()
    }
    candidate = 10
    while candidate in used_codes:
        candidate += 1
    return str(candidate)


def next_act_number(
    con,
    engineer_code,
    now,
    status="completed",
    number_type="standard",
):
    month = now.strftime("%m")

    if status == "draft" and number_type == "thermo":
        prefix = f"DT-{engineer_code}-{month}-"
    elif status == "draft":
        prefix = f"D-{engineer_code}-{month}-"
    elif number_type == "thermo":
        prefix = f"T-{engineer_code}-{month}-"
    else:
        prefix = f"{engineer_code}-{month}-"

    row = con.execute(
        """
        SELECT act_number
        FROM acts
        WHERE act_number LIKE ?
        ORDER BY act_number DESC
        LIMIT 1
        """,
        (prefix + "%",),
    ).fetchone()

    sequence = (
        int(row["act_number"].rsplit("-", 1)[-1]) + 1
        if row
        else 1
    )

    return f"{prefix}{sequence:04d}"


def import_google_sheet():
    request_obj = urllib.request.Request(
        SHEET_URL,
        headers={"User-Agent": "RepairActs/1.0"},
    )

    with urllib.request.urlopen(request_obj, timeout=30) as response:
        content = response.read().decode("utf-8-sig")

    reader = csv.DictReader(io.StringIO(content))
    if not reader.fieldnames:
        raise ValueError("В Google Sheets не найдена строка заголовков.")

    now = datetime.now().isoformat(timespec="seconds")
    imported = 0

    with db() as con:
        for line_number, row in enumerate(reader, start=2):
            sap_id = field(row, "SAP")
            mvz = field(row, "МВЗ", "MB3", "МВЗ ")
            customer = field(row, "клиент", "Клиент")
            object_format = field(row, "Формат объект", "Формат объекта")
            address = field(row, "Адрес объекта")
            model = field(row, "Модель принтера", "Модель принт", "Модель")
            serial = field(row, "Серийный номер", "Серийный ном", "Серийник")

            if not serial or not customer:
                continue

            con.execute(
                """
                INSERT INTO devices (
                    source, source_key, sap_id, mvz, customer_name,
                    object_format, address, model, serial_number,
                    serial_normalized, imported_at
                ) VALUES (?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?)
                ON CONFLICT(source, source_key) DO UPDATE SET
                    sap_id=excluded.sap_id,
                    mvz=excluded.mvz,
                    customer_name=excluded.customer_name,
                    object_format=excluded.object_format,
                    address=excluded.address,
                    model=excluded.model,
                    serial_number=excluded.serial_number,
                    serial_normalized=excluded.serial_normalized,
                    imported_at=excluded.imported_at
                """,
                (
                    "google_sheet_main",
                    str(line_number),
                    sap_id,
                    mvz,
                    customer,
                    object_format,
                    address,
                    model,
                    serial,
                    normalize_serial(serial),
                    now,
                ),
            )
            imported += 1

    return imported


def save_manual_device(con, values, now):
    serial = clean(values.get("serial_number"))
    customer = clean(values.get("customer_name"))

    if not serial or not customer:
        return None

    serial_normalized = normalize_serial(serial)
    existing = con.execute(
        """
        SELECT id
        FROM devices
        WHERE source = 'manual' AND serial_normalized = ?
        LIMIT 1
        """,
        (serial_normalized,),
    ).fetchone()

    data = (
        customer,
        clean(values.get("address")),
        clean(values.get("device_model")),
        serial,
        serial_normalized,
        now.isoformat(timespec="seconds"),
    )

    if existing:
        con.execute(
            """
            UPDATE devices SET
                customer_name=?,
                address=?,
                model=?,
                serial_number=?,
                serial_normalized=?,
                imported_at=?
            WHERE id=?
            """,
            data + (existing["id"],),
        )
        return existing["id"]

    cursor = con.execute(
        """
        INSERT INTO devices (
            source, source_key, sap_id, mvz, customer_name,
            object_format, address, model, serial_number,
            serial_normalized, imported_at
        ) VALUES (?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?)
        """,
        (
            "manual",
            f"manual:{serial_normalized}",
            "",
            "",
            customer,
            "",
            clean(values.get("address")),
            clean(values.get("device_model")),
            serial,
            serial_normalized,
            now.isoformat(timespec="seconds"),
        ),
    )
    return cursor.lastrowid


def form_values():
    names = [
        "customer_name",
        "customer_representative",
        "device_model",
        "serial_number",
        "printeco_label",
        "device_type",
        "comment_label",
        "address",
        "phone",
        "fault",
        "counter_bw",
        "counter_color",
        "service_kind",
        "diagnostics_result",
        "works_text",
        "materials_text",
        "work_date",
        "start_time",
        "end_time",
        "device_condition",
        "customer_signatory",
        "intraservice_task_id",
    ]

    values = {}
    for name in names:
        value = request.form.get(name, "")
        if name == "materials_text":
            values[name] = clean_multiline(value)
        else:
            values[name] = clean(value)
    return values


@app.route("/")
def home():
    engineer = get_engineer()
    if not engineer:
        return redirect(url_for("choose_engineer"))
    return render_template("home.html", engineer=engineer)


@app.route("/engineer", methods=["GET", "POST"])
def choose_engineer():
    if request.method == "POST":
        flash(
            "Вход через эту страницу отключён. "
            "Используйте веб-приложение.",
            "danger",
        )
    return render_template("engineer.html", engineers=get_engineers_dict())


@app.route("/logout")
def logout():
    session.clear()
    return redirect(url_for("choose_engineer"))


@app.route("/api/session", methods=["GET"])
def api_session():
    engineer = get_engineer()
    if not engineer:
        return jsonify({"engineer": None})

    return jsonify({
        "engineer": {
            "id": session.get("engineer_key"),
            "name": engineer["name"],
            "first_name": engineer["first_name"],
            "code": engineer["code"],
            "is_admin": engineer["is_admin"],
        }
    })


@app.route("/api/session", methods=["DELETE"])
def api_session_logout():
    session.clear()
    return jsonify({"engineer": None})


@app.route("/api/session/claim", methods=["POST"])
def api_session_claim():
    data = request.get_json(silent=True) or {}
    key = data.get("engineer_key", "")
    first_name = clean(data.get("first_name", ""))
    password = data.get("password", "")

    if not first_name:
        return jsonify({
            "error": "Укажите имя",
            "code": "MISSING_FIELDS",
        }), 400

    if len(password) < 6:
        return jsonify({
            "error": "Пароль должен быть не короче 6 символов",
            "code": "WEAK_PASSWORD",
        }), 400

    with db() as con:
        row = con.execute(
            "SELECT key, name, code, password_hash, is_admin FROM engineers "
            "WHERE key = ? AND status = 'active'",
            (key,),
        ).fetchone()

        if not row:
            return jsonify({
                "error": "Инженер не найден",
                "code": "NOT_FOUND",
            }), 404

        if row["password_hash"] is not None:
            return jsonify({
                "error": "Пароль уже задан, используйте вход",
                "code": "ALREADY_CLAIMED",
            }), 400

        con.execute(
            "UPDATE engineers SET password_hash = ?, first_name = ? "
            "WHERE key = ?",
            (generate_password_hash(password), first_name, key),
        )

    session["engineer_key"] = key
    return jsonify({
        "engineer": {
            "id": row["key"],
            "name": row["name"],
            "first_name": first_name,
            "code": row["code"],
            "is_admin": bool(row["is_admin"]),
        }
    })


@app.route("/api/session/login", methods=["POST"])
def api_session_login():
    data = request.get_json(silent=True) or {}
    key = data.get("engineer_key", "")
    password = data.get("password", "")

    with db() as con:
        row = con.execute(
            "SELECT key, name, first_name, code, password_hash, is_admin "
            "FROM engineers WHERE key = ? AND status = 'active'",
            (key,),
        ).fetchone()

    if not row or not row["password_hash"]:
        return jsonify({
            "error": "Неверный логин или пароль",
            "code": "INVALID_CREDENTIALS",
        }), 401

    if not check_password_hash(row["password_hash"], password):
        return jsonify({
            "error": "Неверный логин или пароль",
            "code": "INVALID_CREDENTIALS",
        }), 401

    session["engineer_key"] = key
    return jsonify({
        "engineer": {
            "id": row["key"],
            "name": row["name"],
            "first_name": row["first_name"] or "",
            "code": row["code"],
            "is_admin": bool(row["is_admin"]),
        }
    })


@app.route("/api/session/register", methods=["POST"])
def api_session_register():
    data = request.get_json(silent=True) or {}
    first_name = clean(data.get("first_name", ""))
    last_name = clean(data.get("last_name", ""))
    password = data.get("password", "")

    if not first_name or not last_name:
        return jsonify({
            "error": "Заполните все поля",
            "code": "MISSING_FIELDS",
        }), 400

    if len(password) < 6:
        return jsonify({
            "error": "Пароль должен быть не короче 6 символов",
            "code": "WEAK_PASSWORD",
        }), 400

    with db() as con:
        duplicate = con.execute(
            "SELECT 1 FROM engineers "
            "WHERE lower(first_name) = lower(?) AND lower(name) = lower(?)",
            (first_name, last_name),
        ).fetchone()

        if duplicate:
            return jsonify({
                "error": (
                    f"Инженер «{first_name} {last_name}» уже "
                    "зарегистрирован. Если это другой человек, "
                    "добавьте цифру к фамилии, например "
                    f"«{last_name} 2»."
                ),
                "code": "DUPLICATE_NAME",
            }), 400

        key = generate_engineer_key(con, first_name, last_name)
        code = generate_engineer_code(con)

        con.execute(
            "INSERT INTO engineers "
            "(key, name, first_name, code, password_hash, is_admin, "
            "status, created_at) "
            "VALUES (?, ?, ?, ?, ?, 0, 'pending', ?)",
            (
                key,
                last_name,
                first_name,
                code,
                generate_password_hash(password),
                datetime.now().isoformat(timespec="seconds"),
            ),
        )

    return jsonify({
        "status": "pending",
        "message": "Заявка отправлена. Дождитесь подтверждения администратором.",
    })


@app.route("/api/admin/pending-engineers", methods=["GET"])
def api_admin_pending_engineers():
    if not require_admin():
        return jsonify({
            "error": "Недостаточно прав",
            "code": "FORBIDDEN",
        }), 403

    with db() as con:
        rows = con.execute(
            "SELECT key, name, first_name, code, created_at FROM engineers "
            "WHERE status = 'pending' ORDER BY created_at"
        ).fetchall()

    return jsonify([
        {
            "id": row["key"],
            "name": row["name"],
            "first_name": row["first_name"] or "",
            "code": row["code"],
            "created_at": row["created_at"],
        }
        for row in rows
    ])


@app.route("/api/admin/engineers/<key>/approve", methods=["POST"])
def api_admin_approve_engineer(key):
    if not require_admin():
        return jsonify({
            "error": "Недостаточно прав",
            "code": "FORBIDDEN",
        }), 403

    with db() as con:
        result = con.execute(
            "UPDATE engineers SET status = 'active' "
            "WHERE key = ? AND status = 'pending'",
            (key,),
        )

        if result.rowcount == 0:
            return jsonify({
                "error": "Заявка не найдена",
                "code": "NOT_FOUND",
            }), 404

    return jsonify({"status": "active"})


@app.route("/api/admin/engineers/<key>/reject", methods=["POST"])
def api_admin_reject_engineer(key):
    if not require_admin():
        return jsonify({
            "error": "Недостаточно прав",
            "code": "FORBIDDEN",
        }), 403

    with db() as con:
        result = con.execute(
            "DELETE FROM engineers WHERE key = ? AND status = 'pending'",
            (key,),
        )

        if result.rowcount == 0:
            return jsonify({
                "error": "Заявка не найдена",
                "code": "NOT_FOUND",
            }), 404

    return jsonify({"status": "rejected"})


@app.route("/api/admin/engineers", methods=["GET"])
def api_admin_engineers():
    if not require_admin():
        return jsonify({
            "error": "Недостаточно прав",
            "code": "FORBIDDEN",
        }), 403

    with db() as con:
        rows = con.execute(
            "SELECT key, name, first_name, code, is_admin FROM engineers "
            "WHERE status = 'active' ORDER BY name"
        ).fetchall()

    return jsonify([
        {
            "id": row["key"],
            "name": row["name"],
            "first_name": row["first_name"] or "",
            "code": row["code"],
            "is_admin": bool(row["is_admin"]),
        }
        for row in rows
    ])


@app.route("/api/admin/engineers/<key>", methods=["DELETE"])
def api_admin_delete_engineer(key):
    if not require_admin():
        return jsonify({
            "error": "Недостаточно прав",
            "code": "FORBIDDEN",
        }), 403

    if key == session.get("engineer_key"):
        return jsonify({
            "error": "Нельзя удалить свою учётную запись",
            "code": "CANNOT_DELETE_SELF",
        }), 400

    with db() as con:
        row = con.execute(
            "SELECT is_admin FROM engineers WHERE key = ? AND status = 'active'",
            (key,),
        ).fetchone()

        if not row:
            return jsonify({
                "error": "Инженер не найден",
                "code": "NOT_FOUND",
            }), 404

        if row["is_admin"]:
            admin_count = con.execute(
                "SELECT COUNT(*) AS n FROM engineers "
                "WHERE is_admin = 1 AND status = 'active'"
            ).fetchone()["n"]

            if admin_count <= 1:
                return jsonify({
                    "error": "Нельзя удалить последнего администратора",
                    "code": "LAST_ADMIN",
                }), 400

        # Soft delete: keep the row so its key/code can never be
        # reassigned to someone else later (act numbering depends on
        # a code staying tied to one person forever).
        con.execute(
            "UPDATE engineers SET status = 'deleted' WHERE key = ?",
            (key,),
        )

    return jsonify({"status": "deleted"})


@app.route("/devices")
def devices():
    if not require_engineer():
        return redirect(url_for("choose_engineer"))

    q = clean(request.args.get("q", ""))
    tasks = []
    intraservice_error = None

    with db() as con:
        base_rows = con.execute(
            """
            SELECT *
            FROM devices
            ORDER BY imported_at DESC, serial_number
            LIMIT 1000
            """
        ).fetchall()

    if q:
        q_norm = q.lower()
        device_rows = [
            row
            for row in base_rows
            if (
                q_norm in (row["serial_number"] or "").lower()
                or q_norm in (row["customer_name"] or "").lower()
                or q_norm in (row["model"] or "").lower()
            )
        ]
    else:
        device_rows = base_rows

    if q:
        try:
            tasks = search_intraservice(q)
        except requests.RequestException as exc:
            intraservice_error = str(exc)
        except Exception as exc:
            intraservice_error = str(exc)

    return render_template(
        "devices.html",
        devices=device_rows,
        tasks=tasks,
        intraservice_error=intraservice_error,
        q=q,
        engineer=get_engineer(),
    )


@app.route("/sync-devices", methods=["POST"])
def sync_devices():
    if not require_engineer():
        return redirect(url_for("choose_engineer"))

    try:
        count = import_google_sheet()
        flash(f"Справочник обновлен: обработано строк — {count}.", "success")
    except Exception as exc:
        flash(f"Не удалось загрузить Google Sheets: {exc}", "danger")
    return redirect(request.referrer or url_for("devices"))


@app.route("/acts/new/<int:device_id>")
def new_act(device_id):
    if not require_engineer():
        return redirect(url_for("choose_engineer"))

    with db() as con:
        device = con.execute(
            "SELECT * FROM devices WHERE id = ?",
            (device_id,),
        ).fetchone()

    if not device:
        flash("Аппарат не найден.", "danger")
        return redirect(url_for("devices"))

    return render_template(
        "act_form.html",
        act=None,
        device=device,
        engineer=get_engineer(),
        intraservice_task_id=clean(
            request.args.get("intraservice_task_id", "")
        ),
        today=datetime.now().strftime("%Y-%m-%d"),
    )


@app.route("/acts/new-manual")
def new_act_manual():
    if not require_engineer():
        return redirect(url_for("choose_engineer"))

    return render_template(
        "act_form.html",
        act=None,
        device=None,
        engineer=get_engineer(),
        intraservice_task_id=clean(
            request.args.get("intraservice_task_id", "")
        ),
        today=datetime.now().strftime("%Y-%m-%d"),
    )


@app.route("/api/intraservice/search")
def api_intraservice_search():
    if not require_engineer():
        return {"error": "Не выбран инженер"}, 403

    query = clean(request.args.get("q", ""))
    if not query:
        return {"tasks": []}

    try:
        return {"tasks": search_intraservice(query)}
    except requests.HTTPError as exc:
        return {"error": f"Intraservice HTTP error: {exc}"}, 502
    except requests.RequestException as exc:
        return {"error": f"Intraservice connection error: {exc}"}, 502
    except Exception as exc:
        return {"error": f"Intraservice error: {exc}"}, 500


@app.route("/acts/save", methods=["POST"])
def save_act():
    if not require_engineer():
        return redirect(url_for("choose_engineer"))

    engineer_key = session["engineer_key"]
    engineer = get_engineer()
    now = datetime.now()
    act_id = request.form.get("act_id")
    status = request.form.get("status", "draft")
    number_type = request.form.get("number_type", "standard")

    if status not in {"draft", "completed"}:
        status = "draft"
    if number_type not in NUMBER_TYPES:
        number_type = "standard"

    values = form_values()

    with db() as con:
        form_device_id = request.form.get("source_device_id") or None
        manual_device_id = save_manual_device(con, values, now)
        source_device_id = manual_device_id or form_device_id

        if act_id:
            existing = con.execute(
                "SELECT * FROM acts WHERE id = ?",
                (act_id,),
            ).fetchone()

            if not existing:
                flash("Акт не найден.", "danger")
                return redirect(url_for("acts"))

            manual_number = clean(request.form.get("act_number"))
            new_number = manual_number or existing["act_number"]

            if (
                existing["status"] == "draft"
                and status == "completed"
                and not manual_number
            ):
                new_number = next_act_number(
                    con,
                    engineer["code"],
                    now,
                    "completed",
                    number_type,
                )

            con.execute(
                """
                UPDATE acts SET
                    act_number=?, number_type=?, updated_at=?, status=?,
                    customer_name=?, customer_representative=?, device_model=?,
                    serial_number=?, printeco_label=?, device_type=?,
                    comment_label=?, address=?, phone=?, fault=?,
                    counter_bw=?, counter_color=?, service_kind=?,
                    diagnostics_result=?, works_text=?, materials_text=?,
                    work_date=?, start_time=?, end_time=?, device_condition=?,
                    customer_signatory=?, source_device_id=?,
                    intraservice_task_id=?
                WHERE id=?
                """,
                (
                    new_number,
                    number_type,
                    now.isoformat(timespec="seconds"),
                    status,
                    values["customer_name"],
                    values["customer_representative"],
                    values["device_model"],
                    values["serial_number"],
                    values["printeco_label"],
                    values["device_type"],
                    values["comment_label"],
                    values["address"],
                    values["phone"],
                    values["fault"],
                    values["counter_bw"],
                    values["counter_color"],
                    values["service_kind"],
                    values["diagnostics_result"],
                    values["works_text"],
                    values["materials_text"],
                    values["work_date"],
                    values["start_time"],
                    values["end_time"],
                    values["device_condition"],
                    values["customer_signatory"],
                    source_device_id,
                    values["intraservice_task_id"],
                    act_id,
                ),
            )
            saved_id = int(act_id)
        else:
            manual_number = clean(request.form.get("act_number"))
            number = manual_number or next_act_number(
                con,
                engineer["code"],
                now,
                status,
                number_type,
            )

            cursor = con.execute(
                """
                INSERT INTO acts (
                    act_number, number_type, engineer_key, engineer_name,
                    engineer_code, created_at, updated_at, status,
                    customer_name, customer_representative, device_model,
                    serial_number, printeco_label, device_type, comment_label,
                    address, phone, fault, counter_bw, counter_color,
                    service_kind, diagnostics_result, works_text, materials_text,
                    work_date, start_time, end_time, device_condition,
                    customer_signatory, source_device_id, intraservice_task_id
                ) VALUES (
                    ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?,
                    ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?
                )
                """,
                (
                    number,
                    number_type,
                    engineer_key,
                    engineer["name"],
                    engineer["code"],
                    now.isoformat(timespec="seconds"),
                    now.isoformat(timespec="seconds"),
                    status,
                    values["customer_name"],
                    values["customer_representative"],
                    values["device_model"],
                    values["serial_number"],
                    values["printeco_label"],
                    values["device_type"],
                    values["comment_label"],
                    values["address"],
                    values["phone"],
                    values["fault"],
                    values["counter_bw"],
                    values["counter_color"],
                    values["service_kind"],
                    values["diagnostics_result"],
                    values["works_text"],
                    values["materials_text"],
                    values["work_date"],
                    values["start_time"],
                    values["end_time"],
                    values["device_condition"],
                    values["customer_signatory"],
                    source_device_id,
                    values["intraservice_task_id"],
                ),
            )
            saved_id = cursor.lastrowid

    flash("Акт сохранен.", "success")
    return redirect(url_for("edit_act", act_id=saved_id))


@app.route("/acts")
def acts():
    if not require_engineer():
        return redirect(url_for("choose_engineer"))

    q = clean(request.args.get("q", ""))
    engineer_filter = request.args.get("engineer", "")
    status_filter = request.args.get("status", "")
    date_sort = request.args.get("date_sort", "newest")
    date_from = clean(request.args.get("date_from", ""))
    date_to = clean(request.args.get("date_to", ""))

    if date_sort not in {"newest", "oldest"}:
        date_sort = "newest"

    engineers_dict = get_engineers_dict()

    sql = "SELECT * FROM acts WHERE 1=1"
    params = []

    if engineer_filter in engineers_dict:
        sql += " AND engineer_key = ?"
        params.append(engineer_filter)
    if status_filter in {"draft", "completed"}:
        sql += " AND status = ?"
        params.append(status_filter)
    if date_from:
        sql += " AND date(created_at) >= date(?)"
        params.append(date_from)
    if date_to:
        sql += " AND date(created_at) <= date(?)"
        params.append(date_to)

    direction = "ASC" if date_sort == "oldest" else "DESC"
    sql += f" ORDER BY datetime(created_at) {direction}, id {direction} LIMIT 200"

    with db() as con:
        base_rows = con.execute(sql, params).fetchall()

    if q:
        q_norm = q.lower()
        rows = [
            row
            for row in base_rows
            if (
                q_norm in (row["act_number"] or "").lower()
                or q_norm in (row["serial_number"] or "").lower()
                or q_norm in (row["customer_name"] or "").lower()
                or q_norm in (row["device_model"] or "").lower()
            )
        ]
    else:
        rows = base_rows

    return render_template(
        "acts.html",
        acts=rows,
        q=q,
        engineers=engineers_dict,
        engineer_filter=engineer_filter,
        status_filter=status_filter,
        date_sort=date_sort,
        date_from=date_from,
        date_to=date_to,
        engineer=get_engineer(),
    )


@app.route("/acts/export.xlsx")
def export_acts():
    if not require_engineer():
        return redirect(url_for("choose_engineer"))

    q = clean(request.args.get("q", ""))
    engineer_filter = request.args.get("engineer", "")
    status_filter = request.args.get("status", "")
    date_sort = request.args.get("date_sort", "newest")
    date_from = clean(request.args.get("date_from", ""))
    date_to = clean(request.args.get("date_to", ""))

    if date_sort not in {"newest", "oldest"}:
        date_sort = "newest"

    sql = """
        SELECT
            serial_number,
            customer_name,
            engineer_name,
            materials_text,
            work_date
        FROM acts
        WHERE 1=1
    """
    params = []

    if engineer_filter in get_engineers_dict():
        sql += " AND engineer_key = ?"
        params.append(engineer_filter)
    if status_filter in {"draft", "completed"}:
        sql += " AND status = ?"
        params.append(status_filter)
    if date_from:
        sql += " AND date(created_at) >= date(?)"
        params.append(date_from)
    if date_to:
        sql += " AND date(created_at) <= date(?)"
        params.append(date_to)

    direction = "ASC" if date_sort == "oldest" else "DESC"
    sql += f" ORDER BY datetime(created_at) {direction}, id {direction}"

    with db() as con:
        rows = con.execute(sql, params).fetchall()

    if q:
        q_lower = q.lower()
        rows = [
            row
            for row in rows
            if (
                q_lower in (row["serial_number"] or "").lower()
                or q_lower in (row["customer_name"] or "").lower()
                or q_lower in (row["engineer_name"] or "").lower()
                or q_lower in (row["materials_text"] or "").lower()
            )
        ]

    wb = Workbook()
    ws = wb.active
    ws.title = "Результаты"

    headers = [
        ("Серийный номер", "serial_number"),
        ("Заказчик", "customer_name"),
        ("Инженер", "engineer_name"),
        ("ЗИП", "materials_text"),
        ("Дата ремонта", "work_date"),
    ]

    ws.append([title for title, _ in headers])

    for cell in ws[1]:
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = PatternFill("solid", fgColor="1F4E78")
        cell.alignment = Alignment(horizontal="center", vertical="center")

    for row in rows:
        ws.append([row[key] or "" for _, key in headers])

    ws.freeze_panes = "A2"
    ws.auto_filter.ref = ws.dimensions

    widths = {
        "serial_number": 22,
        "customer_name": 35,
        "engineer_name": 22,
        "materials_text": 45,
        "work_date": 16,
    }

    for index, (_, key) in enumerate(headers, start=1):
        ws.column_dimensions[get_column_letter(index)].width = widths[key]

    for row in ws.iter_rows(min_row=2):
        for cell in row:
            cell.alignment = Alignment(vertical="top", wrap_text=True)

    output = BytesIO()
    wb.save(output)
    output.seek(0)

    filename = f"rezultaty_aktov_{datetime.now():%Y-%m-%d}.xlsx"

    return send_file(
        output,
        as_attachment=True,
        download_name=filename,
        mimetype=(
            "application/vnd.openxmlformats-officedocument."
            "spreadsheetml.sheet"
        ),
    )


@app.route("/acts/export-short.xlsx")
def export_short_acts():
    if not require_engineer():
        return redirect(url_for("choose_engineer"))

    q = clean(request.args.get("q", ""))
    engineer_filter = request.args.get("engineer", "")
    status_filter = request.args.get("status", "")
    date_sort = request.args.get("date_sort", "newest")
    date_from = clean(request.args.get("date_from", ""))
    date_to = clean(request.args.get("date_to", ""))

    if date_sort not in {"newest", "oldest"}:
        date_sort = "newest"

    sql = "SELECT * FROM acts WHERE 1=1"
    params = []

    if engineer_filter in get_engineers_dict():
        sql += " AND engineer_key = ?"
        params.append(engineer_filter)

    if status_filter in {"draft", "completed"}:
        sql += " AND status = ?"
        params.append(status_filter)

    if date_from:
        sql += " AND date(created_at) >= date(?)"
        params.append(date_from)

    if date_to:
        sql += " AND date(created_at) <= date(?)"
        params.append(date_to)

    direction = "ASC" if date_sort == "oldest" else "DESC"
    sql += f" ORDER BY datetime(created_at) {direction}, id {direction}"

    with db() as con:
        rows = con.execute(sql, params).fetchall()

    if q:
        q_lower = q.lower()
        rows = [
            row
            for row in rows
            if (
                q_lower in (row["act_number"] or "").lower()
                or q_lower in (row["serial_number"] or "").lower()
                or q_lower in (row["customer_name"] or "").lower()
                or q_lower in (row["device_model"] or "").lower()
            )
        ]

    workbook = Workbook()
    worksheet = workbook.active
    worksheet.title = "Краткая выгрузка"

    worksheet.append(["Заказчик", "Серийный номер"])

    for cell in worksheet[1]:
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = PatternFill("solid", fgColor="1F4E78")
        cell.alignment = Alignment(horizontal="center", vertical="center")

    for row in rows:
        worksheet.append([
            row["customer_name"] or "",
            row["serial_number"] or "",
        ])

    worksheet.column_dimensions["A"].width = 40
    worksheet.column_dimensions["B"].width = 25
    worksheet.freeze_panes = "A2"
    worksheet.auto_filter.ref = worksheet.dimensions

    for row in worksheet.iter_rows(min_row=2):
        for cell in row:
            cell.alignment = Alignment(vertical="top", wrap_text=True)

    output = BytesIO()
    workbook.save(output)
    output.seek(0)

    filename = f"kratkaya_vygruzka_{datetime.now():%Y-%m-%d}.xlsx"

    return send_file(
        output,
        as_attachment=True,
        download_name=filename,
        mimetype=(
            "application/vnd.openxmlformats-officedocument."
            "spreadsheetml.sheet"
        ),
    )


@app.route("/acts/<int:act_id>")
def edit_act(act_id):
    if not require_engineer():
        return redirect(url_for("choose_engineer"))

    with db() as con:
        act = con.execute(
            "SELECT * FROM acts WHERE id = ?",
            (act_id,),
        ).fetchone()

    if not act:
        flash("Акт не найден.", "danger")
        return redirect(url_for("acts"))

    return render_template(
        "act_form.html",
        act=act,
        device=None,
        engineer=get_engineer(),
        intraservice_task_id=act["intraservice_task_id"] or "",
        today=datetime.now().strftime("%Y-%m-%d"),
    )


@app.route("/acts/<int:act_id>/print")
def print_act(act_id):
    if not require_engineer():
        return redirect(url_for("choose_engineer"))

    with db() as con:
        act = con.execute(
            "SELECT * FROM acts WHERE id = ?",
            (act_id,),
        ).fetchone()

    if not act:
        flash("Акт не найден.", "danger")
        return redirect(url_for("acts"))

    return render_template("act_print.html", act=act)


@app.route("/acts/<int:act_id>/delete", methods=["POST"])
def delete_act_web(act_id):
    if not require_engineer():
        return redirect(url_for("choose_engineer"))

    with db() as con:
        act = con.execute(
            "SELECT id, act_number FROM acts WHERE id = ?",
            (act_id,),
        ).fetchone()
        if not act:
            flash("Акт не найден.", "danger")
            return redirect(url_for("acts"))
        con.execute("DELETE FROM acts WHERE id = ?", (act_id,))

    flash(f"Акт {act['act_number']} удален.", "success")
    return redirect(url_for("acts"))


@app.route("/acts/<int:act_id>/renumber", methods=["POST"])
def renumber_act(act_id):
    if not require_engineer():
        return redirect(url_for("choose_engineer"))

    new_number = clean(request.form.get("new_number"))
    if not new_number:
        flash("Введите новый номер.", "danger")
        return redirect(url_for("edit_act", act_id=act_id))

    with db() as con:
        act = con.execute(
            "SELECT id, act_number FROM acts WHERE id = ?",
            (act_id,),
        ).fetchone()
        if not act:
            flash("Акт не найден.", "danger")
            return redirect(url_for("acts"))

        duplicate = con.execute(
            "SELECT id FROM acts WHERE act_number = ? AND id <> ?",
            (new_number, act_id),
        ).fetchone()
        if duplicate:
            flash("Такой номер уже используется.", "danger")
            return redirect(url_for("edit_act", act_id=act_id))

        con.execute(
            """
            UPDATE acts
            SET act_number=?, updated_at=?
            WHERE id=?
            """,
            (
                new_number,
                datetime.now().isoformat(timespec="seconds"),
                act_id,
            ),
        )

    flash(f"Номер изменён: {act['act_number']} → {new_number}.", "success")
    return redirect(url_for("edit_act", act_id=act_id))


@app.route("/acts/<int:act_id>/export.xlsx")
def export_filled_act(act_id):
    if not require_engineer():
        return redirect(url_for("choose_engineer"))

    with db() as con:
        act = con.execute(
            "SELECT * FROM acts WHERE id = ?",
            (act_id,),
        ).fetchone()

    if not act:
        flash("Акт не найден.", "danger")
        return redirect(url_for("acts"))

    wb = Workbook()
    ws = wb.active
    ws.title = "Акт"

    rows = [
        ("Внутренний номер", act["act_number"]),
        (
            "Тип нумерации",
            NUMBER_TYPES.get(
                act["number_type"],
                NUMBER_TYPES["standard"],
            )["name"],
        ),
        (
            "Статус",
            "Черновик" if act["status"] == "draft" else "Завершён",
        ),
        ("Заказчик", act["customer_name"]),
        ("Инженер", act["engineer_name"]),
        ("Модель аппарата", act["device_model"]),
        ("Серийный номер", act["serial_number"]),
        ("Маркировка ПРИНТЭКО", act["printeco_label"]),
        ("Тип аппарата", act["device_type"]),
        ("Счётчик ч/б", act["counter_bw"]),
        ("Счётчик цветной", act["counter_color"]),
        ("Вид услуги", act["service_kind"]),
        ("Результат диагностики", act["diagnostics_result"]),
        ("Выполненные работы", act["works_text"]),
        ("ЗИП / материалы", act["materials_text"]),
        ("Дата ремонта", act["work_date"]),
        ("Время начала", act["start_time"]),
        ("Время окончания", act["end_time"]),
        ("Состояние аппарата", act["device_condition"]),
        ("Подписант заказчика", act["customer_signatory"]),
        ("Создан", act["created_at"]),
        ("Изменён", act["updated_at"]),
    ]

    ws.append(["Поле", "Значение"])
    for key, value in rows:
        ws.append([key, value or ""])

    for cell in ws[1]:
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = PatternFill("solid", fgColor="1F4E78")
        cell.alignment = Alignment(horizontal="center", vertical="center")

    for cell in ws["A"][1:]:
        cell.font = Font(bold=True)
        cell.fill = PatternFill("solid", fgColor="D9EAF7")

    ws.column_dimensions["A"].width = 32
    ws.column_dimensions["B"].width = 70
    ws.freeze_panes = "A2"
    ws.auto_filter.ref = ws.dimensions

    for row in ws.iter_rows(min_row=2):
        for cell in row:
            cell.alignment = Alignment(vertical="top", wrap_text=True)

    output = BytesIO()
    wb.save(output)
    output.seek(0)

    filename = f"act_{clean(act['act_number']) or act_id}.xlsx"

    return send_file(
        output,
        as_attachment=True,
        download_name=filename,
        mimetype=(
            "application/vnd.openxmlformats-officedocument."
            "spreadsheetml.sheet"
        ),
    )


@app.route("/api/engineers", methods=["GET"])
def api_engineers():
    with db() as con:
        rows = con.execute(
            "SELECT key, name, first_name, code, password_hash "
            "FROM engineers WHERE status = 'active' ORDER BY name"
        ).fetchall()

    return jsonify([
        {
            "id": row["key"],
            "name": row["name"],
            "first_name": row["first_name"] or "",
            "code": row["code"],
            "has_password": row["password_hash"] is not None,
        }
        for row in rows
    ])


@app.route("/api/devices", methods=["GET"])
def api_devices():
    if not get_engineer():
        return jsonify({
            "error": "Не выбран инженер",
            "code": "UNAUTHENTICATED",
        }), 401

    q = clean(request.args.get("q", ""))

    if not q:
        return jsonify([])

    q_norm = normalize_serial(q)

    with db() as con:
        rows = con.execute(
            """
            SELECT
                id,
                customer_name,
                model,
                serial_number,
                address
            FROM devices
            WHERE serial_normalized LIKE ?
            ORDER BY customer_name, model
            LIMIT 100
            """,
            (f"%{q_norm}%",),
        ).fetchall()

    return jsonify([
        {
            "id": row["id"],
            "customer_name": row["customer_name"] or "",
            "device_model": row["model"] or "",
            "serial_number": row["serial_number"] or "",
            "address": row["address"] or "",
        }
        for row in rows
    ])


@app.route("/api/acts", methods=["GET"])
def api_acts_list():
    if not get_engineer():
        return jsonify({
            "error": "Не выбран инженер",
            "code": "UNAUTHENTICATED",
        }), 401

    q = clean(request.args.get("q", ""))
    engineer_filter = clean(request.args.get("engineer", ""))
    status_filter = clean(request.args.get("status", ""))
    date_from = clean(request.args.get("date_from", ""))
    date_to = clean(request.args.get("date_to", ""))

    sql = """
        SELECT
            id,
            act_number,
            status,
            customer_name,
            device_model,
            serial_number,
            engineer_name,
            engineer_key,
            created_at,
            work_date
        FROM acts
        WHERE 1=1
    """
    params = []

    if q:
        sql += """
            AND (
                act_number LIKE ?
                OR serial_number LIKE ?
                OR customer_name LIKE ?
                OR device_model LIKE ?
            )
        """
        params.extend([f"%{q}%"] * 4)

    if engineer_filter in get_engineers_dict():
        sql += " AND engineer_key = ?"
        params.append(engineer_filter)

    if status_filter in {"draft", "completed"}:
        sql += " AND status = ?"
        params.append(status_filter)

    if date_from:
        sql += " AND DATE(created_at) >= DATE(?)"
        params.append(date_from)

    if date_to:
        sql += " AND DATE(created_at) <= DATE(?)"
        params.append(date_to)

    sql += " ORDER BY datetime(created_at) DESC, id DESC LIMIT 200"

    with db() as con:
        rows = con.execute(sql, params).fetchall()

    return jsonify([
        {
            "id": row["id"],
            "act_number": row["act_number"] or "",
            "status": row["status"],
            "customer_name": row["customer_name"] or "",
            "device_model": row["device_model"] or "",
            "serial_number": row["serial_number"] or "",
            "engineer_name": row["engineer_name"] or "",
            "engineer_key": row["engineer_key"] or "",
            "created_at": row["created_at"] or "",
            "work_date": row["work_date"] or "",
        }
        for row in rows
    ])


@app.route("/api/acts/<int:act_id>", methods=["GET"])
def api_act_detail(act_id):
    if not get_engineer():
        return jsonify({
            "error": "Не выбран инженер",
            "code": "UNAUTHENTICATED",
        }), 401

    with db() as con:
        act = con.execute(
            "SELECT * FROM acts WHERE id = ?",
            (act_id,),
        ).fetchone()

    if not act:
        return jsonify({
            "error": "Акт не найден",
            "code": "NOT_FOUND",
        }), 404

    return jsonify({
        "id": act["id"],
        "act_number": act["act_number"] or "",
        "number_type": act["number_type"] or "standard",
        "status": act["status"] or "",
        "engineer_key": act["engineer_key"] or "",
        "engineer_name": act["engineer_name"] or "",
        "engineer_code": act["engineer_code"] or "",
        "customer_name": act["customer_name"] or "",
        "customer_representative": act["customer_representative"] or "",
        "device_model": act["device_model"] or "",
        "serial_number": act["serial_number"] or "",
        "printeco_label": act["printeco_label"] or "",
        "device_type": act["device_type"] or "",
        "comment_label": act["comment_label"] or "",
        "address": act["address"] or "",
        "phone": act["phone"] or "",
        "fault": act["fault"] or "",
        "counter_bw": act["counter_bw"] or "",
        "counter_color": act["counter_color"] or "",
        "service_kind": act["service_kind"] or "",
        "diagnostics_result": act["diagnostics_result"] or "",
        "works_text": act["works_text"] or "",
        "materials_text": act["materials_text"] or "",
        "work_date": act["work_date"] or "",
        "start_time": act["start_time"] or "",
        "end_time": act["end_time"] or "",
        "device_condition": act["device_condition"] or "",
        "customer_signatory": act["customer_signatory"] or "",
        "source_device_id": act["source_device_id"],
        "intraservice_task_id": act["intraservice_task_id"] or "",
        "created_at": act["created_at"] or "",
        "updated_at": act["updated_at"] or "",
    })

if __name__ == "__main__":
    init_db()
    app.run(host="0.0.0.0", port=5000, debug=False)
